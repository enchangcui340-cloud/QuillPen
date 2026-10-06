# -*- coding: utf-8 -*-
"""生成 P4（文档保真读取）测试用的全部夹具。

用法：python make-doc-fixtures.py <输出目录>

产出：
  · 扫描件.pdf     纯图片、**无文本层**的 3 页 PDF（最关键的那类文件）
  · 文字版.pdf     有文本层的 PDF（用 Edge/Chrome 无头打印 HTML 生成；没有浏览器就跳过）
  · 报告.docx      python-docx 写的 Word（含标题/正文/表格/中文）
  · 表格.xlsx      openpyxl 写的 Excel（含中文单元格与数字）
  · 图.png         Pillow 画的图片
"""
import os
import subprocess
import sys

from PIL import Image, ImageDraw, ImageFont

OUT = sys.argv[1] if len(sys.argv) > 1 else r"D:\DSH\test01\.tmp\doc-fixtures"
os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------- 字体
def font(size):
    for p in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\arial.ttf"):
        try:
            return ImageFont.truetype(p, size)
        except OSError:
            pass
    return ImageFont.load_default()

# ---------------------------------------------------------------- 1) 扫描件 PDF（纯图片，无文本层）
PAGES = [
    ("扫描页 1", "关键字：ALPHA-2026", "这是一张纯图片页，没有文本层。"),
    ("扫描页 2", "关键字：BETA-2027", "如果读取时没渲染成图，就会什么都读不到。"),
    ("扫描页 3", "关键字：GAMMA-2028", "第三页用来确认多页都能逐页渲染。"),
]
images = []
for i, (title, key, desc) in enumerate(PAGES, start=1):
    W, H = 1240, 1754
    img = Image.new("RGB", (W, H), (252, 252, 250))
    d = ImageDraw.Draw(img)
    d.rectangle([60, 60, W - 60, H - 60], outline=(180, 180, 180), width=3)
    d.text((100, 120), title, font=font(56), fill=(20, 20, 20))
    d.text((100, 240), key, font=font(72), fill=(10, 10, 10))
    d.text((100, 380), desc, font=font(44), fill=(40, 40, 40))
    for y in range(500, 900, 40):
        d.line([(100, y), (W - 100, y)], fill=(230, 230, 230), width=2)
    d.ellipse([W - 420, H - 420, W - 120, H - 120], outline=(90, 120, 200), width=8)   # 蓝色圆圈：像素校验用
    d.text((W - 400, H - 300), f"P{i}", font=font(90), fill=(90, 120, 200))
    images.append(img)
scan_pdf = os.path.join(OUT, "扫描件.pdf")
images[0].save(scan_pdf, save_all=True, append_images=images[1:], resolution=150.0)
print("扫描件.pdf", os.path.getsize(scan_pdf), "字节（3 页，无文本层）")

# ---------------------------------------------------------------- 2) 文字版 PDF（有文本层）
html = os.path.join(OUT, "_文字版.html")
with open(html, "w", encoding="utf-8") as f:
    f.write("""<!doctype html><html><head><meta charset="utf-8"><style>
    body{font-family:"Microsoft YaHei",sans-serif;font-size:20px}
    h1{font-size:32px}</style></head><body>
    <h1>文字版文档标题</h1>
    <p>这一段带文本层，关键字：TEXT-LAYER-2029。</p>
    <p>抽取文本时应该能读到这一行。</p>
    </body></html>""")
text_pdf = os.path.join(OUT, "文字版.pdf")
browsers = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
]
made = False
for b in browsers:
    if not os.path.exists(b):
        continue
    try:
        subprocess.run(
            [b, "--headless", "--disable-gpu", "--no-pdf-header-footer",
             "--print-to-pdf=" + text_pdf, "file:///" + html.replace("\\", "/")],
            timeout=90, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if os.path.exists(text_pdf) and os.path.getsize(text_pdf) > 1000:
            print("文字版.pdf", os.path.getsize(text_pdf), "字节（有文本层）")
            made = True
            break
    except Exception as e:  # noqa: BLE001
        print("  用", os.path.basename(b), "生成失败：", e)
if not made:
    print("文字版.pdf 跳过（没找到可用的无头浏览器）")

# ---------------------------------------------------------------- 3) docx
try:
    from docx import Document
    doc = Document()
    doc.add_heading("季度报告标题", level=1)
    doc.add_paragraph("这是 Word 正文，关键字：DOCX-MARKER-2030，中文也要能读出来。")
    doc.add_paragraph("第二段：用来确认多段落都抽得到。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "测试项"
    table.cell(1, 1).text = "42"
    docx_path = os.path.join(OUT, "报告.docx")
    doc.save(docx_path)
    print("报告.docx", os.path.getsize(docx_path), "字节")
except Exception as e:  # noqa: BLE001
    print("报告.docx 生成失败：", e)

# ---------------------------------------------------------------- 4) xlsx
try:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"] = "名称"
    ws["B1"] = "数值"
    ws["A2"] = "测试行"
    ws["B2"] = 12345
    ws["A3"] = "关键字XLSX-MARKER-2031"
    xlsx_path = os.path.join(OUT, "表格.xlsx")
    wb.save(xlsx_path)
    print("表格.xlsx", os.path.getsize(xlsx_path), "字节")
except Exception as e:  # noqa: BLE001
    print("表格.xlsx 生成失败：", e)

# ---------------------------------------------------------------- 5) 图片
img = Image.new("RGB", (900, 500), (255, 255, 255))
d = ImageDraw.Draw(img)
d.rectangle([20, 20, 880, 480], outline=(200, 60, 60), width=6)
d.text((60, 100), "PIC-MARKER-2032", font=font(64), fill=(20, 20, 20))
png_path = os.path.join(OUT, "图.png")
img.save(png_path)
print("图.png", os.path.getsize(png_path), "字节")
print("输出目录：", OUT)
